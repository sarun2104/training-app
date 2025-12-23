"""
Export PostgreSQL Database Schema to Excel
Extracts all table names and column information from training_db
"""
import psycopg2
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import os

# Database connection settings
DB_CONFIG = {
    "host": os.getenv("POSTGRES_HOST", "localhost"),
    "port": int(os.getenv("POSTGRES_PORT", 5432)),
    "user": os.getenv("POSTGRES_USER", "postgres"),
    "password": os.getenv("POSTGRES_PASSWORD", "postgres"),
    "database": os.getenv("POSTGRES_DB", "training_db")
}

def get_all_tables_and_columns():
    """Query PostgreSQL information_schema to get all tables and columns"""
    
    query = """
    SELECT 
        t.table_name,
        c.column_name,
        c.data_type,
        c.character_maximum_length,
        c.is_nullable,
        c.column_default,
        c.ordinal_position
    FROM information_schema.tables t
    JOIN information_schema.columns c 
        ON t.table_name = c.table_name 
        AND t.table_schema = c.table_schema
    WHERE t.table_schema = 'public'
        AND t.table_type = 'BASE TABLE'
    ORDER BY t.table_name, c.ordinal_position;
    """
    
    conn = psycopg2.connect(**DB_CONFIG)
    try:
        cursor = conn.cursor()
        cursor.execute(query)
        columns = [desc[0] for desc in cursor.description]
        rows = cursor.fetchall()
        return columns, rows
    finally:
        conn.close()


def create_excel_report(columns, rows, output_path):
    """Create a formatted Excel file with the schema information"""
    
    # Create DataFrame
    df = pd.DataFrame(rows, columns=columns)
    
    # Rename columns for better readability
    df.columns = [
        'Table Name',
        'Column Name', 
        'Data Type',
        'Max Length',
        'Is Nullable',
        'Default Value',
        'Position'
    ]
    
    # Create Excel workbook
    wb = Workbook()
    
    # ==================== Sheet 1: All Columns Detail ====================
    ws1 = wb.active
    ws1.title = "All Columns Detail"
    
    # Add header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write DataFrame to sheet
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border
            if r_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
    
    # Adjust column widths
    ws1.column_dimensions['A'].width = 30
    ws1.column_dimensions['B'].width = 30
    ws1.column_dimensions['C'].width = 25
    ws1.column_dimensions['D'].width = 12
    ws1.column_dimensions['E'].width = 12
    ws1.column_dimensions['F'].width = 30
    ws1.column_dimensions['G'].width = 10
    
    # ==================== Sheet 2: Tables Summary ====================
    ws2 = wb.create_sheet(title="Tables Summary")
    
    # Get unique tables with column counts
    table_summary = df.groupby('Table Name').agg({
        'Column Name': 'count'
    }).reset_index()
    table_summary.columns = ['Table Name', 'Column Count']
    
    # Add columns list for each table
    table_columns = df.groupby('Table Name')['Column Name'].apply(lambda x: ', '.join(x)).reset_index()
    table_columns.columns = ['Table Name', 'Columns']
    table_summary = table_summary.merge(table_columns, on='Table Name')
    
    # Write summary to sheet
    for r_idx, row in enumerate(dataframe_to_rows(table_summary, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border
            if r_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
    
    # Adjust column widths
    ws2.column_dimensions['A'].width = 30
    ws2.column_dimensions['B'].width = 15
    ws2.column_dimensions['C'].width = 100
    
    # ==================== Sheet 3: Individual Table Sheets ====================
    tables = df['Table Name'].unique()
    for table_name in tables:
        # Create sheet for each table (truncate name if too long)
        sheet_name = table_name[:31] if len(table_name) > 31 else table_name
        ws = wb.create_sheet(title=sheet_name)
        
        # Filter data for this table
        table_df = df[df['Table Name'] == table_name][['Column Name', 'Data Type', 'Max Length', 'Is Nullable', 'Default Value']]
        
        # Add table name as header
        ws.cell(row=1, column=1, value=f"Table: {table_name}")
        ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="2F5496")
        ws.merge_cells('A1:E1')
        
        # Write data
        for r_idx, row in enumerate(dataframe_to_rows(table_df, index=False, header=True), 3):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = thin_border
                if r_idx == 3:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center')
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 30
    
    # Save workbook
    wb.save(output_path)
    print(f"Excel file saved to: {output_path}")
    return table_summary


def main():
    print("Connecting to PostgreSQL database: training_db")
    print(f"Host: {DB_CONFIG['host']}:{DB_CONFIG['port']}")
    print("-" * 50)
    
    try:
        columns, rows = get_all_tables_and_columns()
        
        if not rows:
            print("No tables found in the database!")
            return
        
        # Output path
        output_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "training_db_schema.xlsx"
        )
        
        # Create Excel report
        table_summary = create_excel_report(columns, rows, output_path)
        
        print("\n" + "=" * 50)
        print("DATABASE SCHEMA EXPORTED SUCCESSFULLY!")
        print("=" * 50)
        print(f"\nTotal Tables Found: {len(table_summary)}")
        print("\nTables:")
        for _, row in table_summary.iterrows():
            print(f"  - {row['Table Name']} ({row['Column Count']} columns)")
        print(f"\nExcel file: {output_path}")
        
    except psycopg2.OperationalError as e:
        print(f"ERROR: Could not connect to database!")
        print(f"Details: {e}")
        print("\nMake sure PostgreSQL is running and accessible.")
    except Exception as e:
        print(f"ERROR: {e}")
        raise


if __name__ == "__main__":
    main()
